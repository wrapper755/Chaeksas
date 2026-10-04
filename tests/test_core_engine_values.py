"""M3 조각 3b — 「값을 만드는 노드」: 규칙(DMN)·파일 목록·파일 출력·메일/웹훅·이정표.

여기서 거듭 보는 것은 셋이다.

1. **실행 폴더 밖은 막힌다** (ADR-0026) — 그리고 그것은 **오류 경계로 받지 않는다** (그림·설정이
   잘못된 것이라 우회할 일이 아니다). 폴더가 없는 것은 업무 실패라서 경계가 받는다.
2. **보내기 어댑터가 없으면 보내지 않고 실패한다** — 조용히 삼키면 보냈는지 아무도 모른다.
3. **업무 값은 기록에 남지 않는다** (계약 원칙 6) — 받는 사람·파일 경로·판정 결과는 개수·형식만.
"""

from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pytest

from chaeksas.contracts.bpmn_ext import read_process
from chaeksas.contracts.dmn import read_decisions
from chaeksas.core.engine import Engine, Run, RunEnv, State
from chaeksas.core.files import Workspace
from chaeksas.core.run_log import RunLog
from chaeksas.core.senders import RecordingSender

NOW = datetime(2026, 10, 4, 9, 30, tzinfo=UTC)
RUN_ID = "run_20261004_093000_abc123"
EXAMPLES = Path(__file__).resolve().parent.parent / "docs" / "08-business-examples" / "bpmn"

SHELL = """<?xml version="1.0" encoding="UTF-8"?>
<bpmn:definitions xmlns:bpmn="http://www.omg.org/spec/BPMN/20100524/MODEL"
                  xmlns:chk="urn:chaeksas:bpmn:1" id="Defs">
  <bpmn:process id="test.proc" name="시험">
    <bpmn:extensionElements><chk:process>{process_props}</chk:process></bpmn:extensionElements>
    {body}
  </bpmn:process>
</bpmn:definitions>
"""


def make(body: str, *, process_props: str = "{}") -> Any:
    return read_process(SHELL.format(body=body, process_props=process_props))


def flow(id_: str, source: str, target: str) -> str:
    return f'<bpmn:sequenceFlow id="{id_}" sourceRef="{source}" targetRef="{target}" />'


def script(id_: str, body: str, *, data: str = "") -> str:
    association = (
        f'<bpmn:dataOutputAssociation id="DOA_{id_}"><bpmn:targetRef>{data}</bpmn:targetRef>'
        "</bpmn:dataOutputAssociation>"
        if data
        else ""
    )
    return (
        f'<bpmn:scriptTask id="{id_}" scriptFormat="chk-expr">{association}'
        f"<bpmn:script>{body}</bpmn:script></bpmn:scriptTask>"
    )


def data_object(id_: str, props: str) -> str:
    return (
        f'<bpmn:dataObjectReference id="{id_}"><bpmn:extensionElements>'
        f"<chk:dataOutput>{props}</chk:dataOutput></bpmn:extensionElements></bpmn:dataObjectReference>"
    )


def straight(*, middle: str, node_id: str, props: str = "{}") -> Any:
    """시작 → `<middle>` → 끝. `middle`은 노드 XML 조각 하나다."""
    return make(
        '<bpmn:startEvent id="Start_1"/>'
        + middle
        + '<bpmn:endEvent id="End_1"/>'
        + flow("f1", "Start_1", node_id)
        + flow("f2", node_id, "End_1"),
        process_props=props,
    )


def env_for(tmp_path: Path, **extra: Any) -> RunEnv:
    """출력 폴더 하나, 읽기 허용 폴더 하나(`share/`), 시험용 보내기 어댑터."""
    outputs = tmp_path / "outputs"
    share = tmp_path / "share"
    outputs.mkdir(exist_ok=True)
    share.mkdir(exist_ok=True)
    defaults: dict[str, Any] = {
        "workspace": Workspace(output_dir=outputs, readable=(share,)),
        "sender": RecordingSender(),
    }
    return RunEnv(**{**defaults, **extra})


def start(process: Any, env: RunEnv | None = None, **kwargs: Any) -> tuple[Engine, Run]:
    engine = Engine()
    log = RunLog(run_id=RUN_ID)
    return engine, engine.start(process, run_id=RUN_ID, log=log, now=NOW, env=env, **kwargs)


def example_decisions() -> dict[str, Any]:
    found: dict[str, Any] = {}
    for path in sorted(EXAMPLES.glob("*.dmn")):
        found.update(read_decisions(path.read_text(encoding="utf-8")))
    return found


# ─────────────────────────── 규칙 (DMN) ───────────────────────────


def test_a_rule_task_puts_the_decision_into_variables(tmp_path: Path) -> None:
    """예제 FX-02를 그대로 돌린다 — 케이스가 적어 둔 기대값까지 간다."""
    process = read_process((EXAMPLES / "fx02_business_rule.bpmn").read_text(encoding="utf-8"))
    engine, run = start(
        process,
        env_for(tmp_path, decisions=example_decisions()),
        inputs={"금액": 1500000, "등급": "VIP"},
    )
    assert engine.run_until_blocked(run) is State.DONE
    assert run.variables["할인율"] == 0.2
    assert run.variables["사유"] == "VIP 대량 주문"
    assert run.variables["최종가"] == 1200000

    states = [e for e in run.log.events if e.kind == "node_state" and e.node_id == "Task_Rule"]
    assert [e.data["task_type"] for e in states] == ["rule", "rule"]
    # **판정 값은 기록에 없다** (원칙 6).
    assert "VIP 대량 주문" not in str(run.log.events)


def test_a_collect_rule_gives_a_list(tmp_path: Path) -> None:
    rule = '{"decision": "card_policy", "input": {"요일": "요일", "시": "시", "업종": "업종", "금액": "금액"},' \
           ' "output": {"사유": "violation"}}'
    process = make(
        '<bpmn:startEvent id="Start_1"/>'
        + script("Task_Row", "요일 = '토'\n시 = 23\n업종 = '유흥'\n금액 = 600000")
        + f'<bpmn:businessRuleTask id="Task_Rule"><bpmn:extensionElements><chk:rule>{rule}</chk:rule>'
        + "</bpmn:extensionElements></bpmn:businessRuleTask>"
        + '<bpmn:endEvent id="End_1"/>'
        + flow("f1", "Start_1", "Task_Row")
        + flow("f2", "Task_Row", "Task_Rule")
        + flow("f3", "Task_Rule", "End_1")
    )
    engine, run = start(process, env_for(tmp_path, decisions=example_decisions()))
    assert engine.run_until_blocked(run) is State.DONE
    assert run.variables["사유"] == ["주말 사용", "심야 사용", "제한 업종", "1회 한도 초과"]


def test_a_rule_without_its_dmn_stops_the_run(tmp_path: Path) -> None:
    """패키지에 `.dmn`이 없으면 **조용히 지나가지 않는다.**"""
    rule = '{"decision": "없는표", "input": {}, "output": {"답": "x"}}'
    process = straight(
        middle=f'<bpmn:businessRuleTask id="Task_Rule"><bpmn:extensionElements><chk:rule>{rule}</chk:rule>'
        "</bpmn:extensionElements></bpmn:businessRuleTask>",
        node_id="Task_Rule",
    )
    engine, run = start(process, env_for(tmp_path))
    assert engine.run_until_blocked(run) is State.FAILED
    assert run.error is not None and run.error.code == "decision_missing"


# ─────────────────────────── 파일 목록 ───────────────────────────


def file_list_process(props: str) -> Any:
    return straight(
        middle=f'<bpmn:serviceTask id="Task_List"><bpmn:extensionElements><chk:fileList>{props}'
        "</chk:fileList></bpmn:extensionElements></bpmn:serviceTask>",
        node_id="Task_List",
    )


def test_a_file_list_task_collects_paths_in_a_fixed_order(tmp_path: Path) -> None:
    """C14 §파일 목록 — 차례는 `sort`가 정한다 (디스크 순서는 실행마다 다르다, ADR-0026)."""
    share = tmp_path / "share"
    share.mkdir()
    for name in ("c.pdf", "a.pdf", "b.pdf", "메모.txt"):
        (share / name).write_text("x", encoding="utf-8")
    (share / "안쪽").mkdir()
    (share / "안쪽" / "d.pdf").write_text("x", encoding="utf-8")

    props = f'{{"folder": "{share.as_posix()}", "pattern": "*.pdf", "store_as": "파일", "count_as": "건수"}}'
    engine, run = start(file_list_process(props), env_for(tmp_path))
    assert engine.run_until_blocked(run) is State.DONE
    assert run.variables["파일"] == [
        f"{share.as_posix()}/a.pdf",
        f"{share.as_posix()}/b.pdf",
        f"{share.as_posix()}/c.pdf",
    ], "폴더 아래 한 겹만, 이름순으로"
    assert run.variables["건수"] == 3
    # 경로는 기록에 없다 — 개수만.
    assert "a.pdf" not in str(run.log.events)


def test_a_recursive_file_list_goes_down_and_an_empty_one_is_not_a_failure(tmp_path: Path) -> None:
    share = tmp_path / "share"
    share.mkdir()
    (share / "안쪽").mkdir()
    (share / "안쪽" / "d.pdf").write_text("x", encoding="utf-8")

    props = (
        f'{{"folder": "{share.as_posix()}", "pattern": "*.pdf", "recursive": true, "store_as": "파일"}}'
    )
    engine, run = start(file_list_process(props), env_for(tmp_path))
    assert engine.run_until_blocked(run) is State.DONE
    assert run.variables["파일"] == [f"{share.as_posix()}/안쪽/d.pdf"]

    props = f'{{"folder": "{share.as_posix()}", "pattern": "*.xlsx", "store_as": "파일"}}'
    engine, run = start(file_list_process(props), env_for(tmp_path))
    assert engine.run_until_blocked(run) is State.DONE
    assert run.variables["파일"] == [], "맞는 파일이 없는 것은 실패가 아니다"


def test_a_folder_outside_the_allowed_roots_is_denied_and_not_caught(tmp_path: Path) -> None:
    """ADR-0026 — 허용 밖 경로는 **오류 경계로 받지 않는다** (고쳐야 할 설정이다)."""
    outside = tmp_path / "밖"
    outside.mkdir()
    props = f'{{"folder": "{outside.as_posix()}", "store_as": "파일"}}'
    process = make(
        '<bpmn:startEvent id="Start_1"/>'
        + f'<bpmn:serviceTask id="Task_List"><bpmn:extensionElements><chk:fileList>{props}'
        + "</chk:fileList></bpmn:extensionElements></bpmn:serviceTask>"
        + '<bpmn:boundaryEvent id="Bnd_Fail" attachedToRef="Task_List">'
        + "<bpmn:errorEventDefinition/></bpmn:boundaryEvent>"
        + script("Task_Recover", "복구 = 1")
        + '<bpmn:endEvent id="End_1"/><bpmn:endEvent id="End_2"/>'
        + flow("f1", "Start_1", "Task_List")
        + flow("f2", "Task_List", "End_1")
        + flow("f3", "Bnd_Fail", "Task_Recover")
        + flow("f4", "Task_Recover", "End_2")
    )
    engine, run = start(process, env_for(tmp_path))
    assert engine.run_until_blocked(run) is State.FAILED
    assert run.error is not None and run.error.code == "path_denied"
    assert "복구" not in run.variables


def test_a_missing_folder_is_a_task_failure_the_boundary_can_catch(tmp_path: Path) -> None:
    """공유 폴더가 안 붙은 날이 있다 — 그것은 업무 실패라서 흐름으로 받을 수 있다."""
    share = tmp_path / "share"
    share.mkdir()
    props = f'{{"folder": "{(share / "없는폴더").as_posix()}", "store_as": "파일"}}'
    process = make(
        '<bpmn:startEvent id="Start_1"/>'
        + f'<bpmn:serviceTask id="Task_List"><bpmn:extensionElements><chk:fileList>{props}'
        + "</chk:fileList></bpmn:extensionElements></bpmn:serviceTask>"
        + '<bpmn:boundaryEvent id="Bnd_Fail" attachedToRef="Task_List">'
        + "<bpmn:errorEventDefinition/></bpmn:boundaryEvent>"
        + script("Task_Recover", "복구 = error_code")
        + '<bpmn:endEvent id="End_1"/><bpmn:endEvent id="End_2"/>'
        + flow("f1", "Start_1", "Task_List")
        + flow("f2", "Task_List", "End_1")
        + flow("f3", "Bnd_Fail", "Task_Recover")
        + flow("f4", "Task_Recover", "End_2")
    )
    engine, run = start(process, env_for(tmp_path))
    assert engine.run_until_blocked(run) is State.DONE
    assert run.variables["복구"] == "TASK_FAILED"


def test_a_relative_folder_is_read_under_the_output_folder(tmp_path: Path) -> None:
    outputs = tmp_path / "outputs"
    outputs.mkdir()
    (outputs / "모음").mkdir()
    (outputs / "모음" / "a.md").write_text("x", encoding="utf-8")
    props = '{"folder": "모음", "store_as": "파일"}'
    engine, run = start(file_list_process(props), env_for(tmp_path))
    assert engine.run_until_blocked(run) is State.DONE
    assert run.variables["파일"] == ["모음/a.md"]


# ─────────────────────────── 파일 출력 ───────────────────────────


def writing_process(props: str, *, script_body: str = "본문 = '# 보고'") -> Any:
    return make(
        '<bpmn:startEvent id="Start_1"/>'
        + script("Task_Make", script_body, data="Data_Out")
        + data_object("Data_Out", props)
        + '<bpmn:endEvent id="End_1"/>'
        + flow("f1", "Start_1", "Task_Make")
        + flow("f2", "Task_Make", "End_1")
    )


def test_a_markdown_file_is_written_when_the_task_ends(tmp_path: Path) -> None:
    props = '{"path": "보고/{오늘}.md", "format": "md", "template": "{본문}", "store_as": "보고경로"}'
    engine, run = start(writing_process(props), env_for(tmp_path))
    assert engine.run_until_blocked(run) is State.DONE
    # `store_as`에는 **BPM 프로세스가 적은 경로 그대로**가 들어간다 (PC의 절대 경로가 아니다).
    assert run.variables["보고경로"] == "보고/2026-10-04.md"
    assert (tmp_path / "outputs" / "보고" / "2026-10-04.md").read_text(encoding="utf-8") == "# 보고\n"
    written = [e for e in run.log.events if "파일을 썼다" in str(e.data.get("message"))]
    assert len(written) == 1 and written[0].node_id == "Data_Out"
    assert "보고/2026-10-04.md" not in str(run.log.events), "경로는 기록에 남기지 않는다"


def test_json_and_csv_and_table_shapes(tmp_path: Path) -> None:
    rows = "표 = [dict(이름='가', 수=1), dict(이름='나', 수=2)]"
    props = '{"path": "결과.json", "format": "json", "variables": ["표", "메모"], "store_as": "경로"}'
    engine, run = start(writing_process(props, script_body=f"{rows}\n메모 = '좋음'"), env_for(tmp_path))
    assert engine.run_until_blocked(run) is State.DONE
    import json  # noqa: PLC0415

    assert json.loads((tmp_path / "outputs" / "결과.json").read_text(encoding="utf-8")) == {
        "표": [{"이름": "가", "수": 1}, {"이름": "나", "수": 2}],
        "메모": "좋음",
    }

    props = '{"path": "결과.csv", "format": "csv", "variables": ["표"], "store_as": "경로"}'
    engine, run = start(writing_process(props, script_body=rows), env_for(tmp_path))
    assert engine.run_until_blocked(run) is State.DONE
    text = (tmp_path / "outputs" / "결과.csv").read_text(encoding="utf-8")
    assert text.startswith("﻿"), "Excel이 UTF-8로 열게 BOM을 붙인다"
    assert text.splitlines()[1:] == ["가,1", "나,2"]

    props = '{"path": "표.md", "format": "md", "variables": ["표"], "store_as": "경로"}'
    engine, run = start(writing_process(props, script_body=rows), env_for(tmp_path))
    assert engine.run_until_blocked(run) is State.DONE
    assert (tmp_path / "outputs" / "표.md").read_text(encoding="utf-8").splitlines() == [
        "## 표", "", "| 이름 | 수 |", "| --- | --- |", "| 가 | 1 |", "| 나 | 2 |",
    ]


def test_appending_does_not_repeat_the_header(tmp_path: Path) -> None:
    props = '{"path": "쌓기.csv", "format": "csv", "variables": ["표"], "append": true, "store_as": "경로"}'
    env = env_for(tmp_path)
    for _ in range(2):
        engine, run = start(writing_process(props, script_body="표 = [dict(이름='가')]"), env)
        assert engine.run_until_blocked(run) is State.DONE
    assert (tmp_path / "outputs" / "쌓기.csv").read_text(encoding="utf-8").splitlines() == [
        "﻿이름", "가", "가",
    ]


def test_an_xlsx_file_gets_one_sheet_per_variable(tmp_path: Path) -> None:
    from openpyxl import load_workbook  # noqa: PLC0415

    props = '{"path": "결과.xlsx", "format": "xlsx", "variables": ["전체", "요약"], "store_as": "경로"}'
    body = "전체 = [dict(이름='가', 수=1)]\n요약 = '한 건'"
    engine, run = start(writing_process(props, script_body=body), env_for(tmp_path))
    assert engine.run_until_blocked(run) is State.DONE

    book = load_workbook(tmp_path / "outputs" / "결과.xlsx")
    assert book.sheetnames == ["전체", "요약"]
    assert [list(r) for r in book["전체"].values] == [["이름", "수"], ["가", 1]]
    assert [list(r) for r in book["요약"].values] == [["요약", "한 건"]]


def test_writing_outside_the_output_folder_is_denied(tmp_path: Path) -> None:
    props = '{"path": "../밖/몰래.md", "format": "md", "template": "{본문}", "store_as": "경로"}'
    engine, run = start(writing_process(props), env_for(tmp_path))
    assert engine.run_until_blocked(run) is State.FAILED
    assert run.error is not None and run.error.code == "path_denied"
    assert not (tmp_path / "밖").exists()


def test_without_an_output_folder_nothing_is_written(tmp_path: Path) -> None:
    """기본 `RunEnv`는 아무것도 못 쓴다 — **조용히 아무 데나 쓰지 않는다.**"""
    props = '{"path": "보고.md", "format": "md", "template": "{본문}", "store_as": "경로"}'
    engine, run = start(writing_process(props))
    assert engine.run_until_blocked(run) is State.FAILED
    assert run.error is not None and run.error.code == "path_denied"


def test_a_file_written_after_an_approval_sees_the_form_answer(tmp_path: Path) -> None:
    """결재 폼의 칸을 그 결재의 출력 파일에 넣는 것이 흔하다 (B11이 같은 가정을 쓴다)."""
    approval = '{"title": "확인", "fields": [{"key": "의견", "label": "의견", "type": "text", "required": true}]}'
    props = '{"path": "결재.md", "format": "md", "template": "의견: {의견}", "store_as": "경로"}'
    process = make(
        '<bpmn:startEvent id="Start_1"/>'
        + '<bpmn:userTask id="Approve_1"><bpmn:extensionElements>'
        + f"<chk:approval>{approval}</chk:approval></bpmn:extensionElements>"
        + '<bpmn:dataOutputAssociation id="DOA_1"><bpmn:targetRef>Data_Out</bpmn:targetRef>'
        + "</bpmn:dataOutputAssociation></bpmn:userTask>"
        + data_object("Data_Out", props)
        + '<bpmn:endEvent id="End_1"/>'
        + flow("f1", "Start_1", "Approve_1")
        + flow("f2", "Approve_1", "End_1")
    )
    engine, run = start(process, env_for(tmp_path))
    assert engine.run_until_blocked(run) is State.WAITING
    assert run.pending is not None
    engine.answer(run, run.pending.request_id, {"의견": "좋습니다"}, answered_by="홍길동")
    assert engine.run_until_blocked(run) is State.DONE
    assert (tmp_path / "outputs" / "결재.md").read_text(encoding="utf-8") == "의견: 좋습니다\n"


# ─────────────────────────── 메일·웹훅 ───────────────────────────


def send_process(element: str, props: str, *, data: str = "") -> Any:
    middle = (
        f'<bpmn:sendTask id="Task_Send"><bpmn:extensionElements>'
        f"<chk:{element}>{props}</chk:{element}></bpmn:extensionElements></bpmn:sendTask>"
    )
    return make(
        '<bpmn:startEvent id="Start_1"/>'
        + script("Task_Make", "본문 = '안녕하세요'\n받는사람 = 'a@example.com'\n건수 = 3", data=data)
        + (data_object("Data_Out", data and '{"path": "보고.md", "format": "md", '
                                            '"template": "{본문}", "store_as": "보고경로"}') if data else "")
        + middle
        + '<bpmn:endEvent id="End_1"/>'
        + flow("f1", "Start_1", "Task_Make")
        + flow("f2", "Task_Make", "Task_Send")
        + flow("f3", "Task_Send", "End_1")
    )


def test_an_email_goes_to_the_adapter_with_its_attachment_resolved(tmp_path: Path) -> None:
    props = (
        '{"to": ["{받는사람}"], "subject": "[시험] {오늘}", "body": "{본문}",'
        ' "attachments": ["보고경로"], "store_as": "발송결과"}'
    )
    env = env_for(tmp_path)
    engine, run = start(send_process("email", props, data="Data_Out"), env)
    assert engine.run_until_blocked(run) is State.DONE

    sender = env.sender
    assert isinstance(sender, RecordingSender)
    (sent,) = sender.emails
    assert sent.to == ("a@example.com",)
    assert sent.subject == "[시험] 2026-10-04"
    assert sent.body == "안녕하세요"
    # 첨부는 **푼 경로**로 어댑터에 간다 (변수에는 적은 경로 그대로 남는다).
    assert sent.attachments == (tmp_path / "outputs" / "보고.md",)
    assert run.variables["발송결과"] == {"ok": True, "to": 1, "id": "rec_1"}
    assert "a@example.com" not in str(run.log.events), "받는 사람은 기록에 남기지 않는다"


def test_without_a_sender_the_mail_is_not_sent_and_the_boundary_catches_it(tmp_path: Path) -> None:
    """기본은 **보내지 않고 실패**다 — 조용히 삼키면 보냈는지 아무도 모른다."""
    props = '{"to": ["{받는사람}"], "subject": "x", "body": "{본문}"}'
    middle = (
        f'<bpmn:sendTask id="Task_Send"><bpmn:extensionElements><chk:email>{props}</chk:email>'
        "</bpmn:extensionElements></bpmn:sendTask>"
    )
    process = make(
        '<bpmn:startEvent id="Start_1"/>'
        + script("Task_Make", "본문 = 'ㅇ'\n받는사람 = 'a@example.com'")
        + middle
        + '<bpmn:boundaryEvent id="Bnd_Fail" attachedToRef="Task_Send">'
        + "<bpmn:errorEventDefinition/></bpmn:boundaryEvent>"
        + script("Task_Recover", "실패 = error_code")
        + '<bpmn:endEvent id="End_1"/><bpmn:endEvent id="End_2"/>'
        + flow("f1", "Start_1", "Task_Make")
        + flow("f2", "Task_Make", "Task_Send")
        + flow("f3", "Task_Send", "End_1")
        + flow("f4", "Bnd_Fail", "Task_Recover")
        + flow("f5", "Task_Recover", "End_2")
    )
    engine, run = start(process, RunEnv(workspace=Workspace(output_dir=tmp_path)))
    assert engine.run_until_blocked(run) is State.DONE
    assert run.variables["실패"] == "SEND_FAILED"


def test_a_webhook_sends_only_the_listed_fields(tmp_path: Path) -> None:
    props = '{"url": "https://example.test/hook", "method": "POST", "body": "fields:[건수]", "store_as": "응답"}'
    env = env_for(tmp_path)
    engine, run = start(send_process("webhook", props), env)
    assert engine.run_until_blocked(run) is State.DONE

    sender = env.sender
    assert isinstance(sender, RecordingSender)
    (sent,) = sender.webhooks
    assert sent.url == "https://example.test/hook"
    assert sent.body == {"건수": 3}, "적은 변수만 간다 (비밀이 섞이지 않게)"
    assert run.variables["응답"] == {"ok": True, "status": 200, "body": ""}


def test_a_webhook_can_send_a_template_or_everything(tmp_path: Path) -> None:
    props = '{"url": "https://example.test/h", "body": "template:\\"건수 {건수}\\""}'
    env = env_for(tmp_path)
    engine, run = start(send_process("webhook", props), env)
    assert engine.run_until_blocked(run) is State.DONE
    sender = env.sender
    assert isinstance(sender, RecordingSender)
    assert sender.webhooks[0].body == "건수 3"

    props = '{"url": "https://example.test/h", "body": "all"}'
    env = env_for(tmp_path)
    engine, run = start(send_process("webhook", props), env)
    assert engine.run_until_blocked(run) is State.DONE
    sender = env.sender
    assert isinstance(sender, RecordingSender)
    assert isinstance(sender.webhooks[0].body, dict)
    assert sender.webhooks[0].body["건수"] == 3


def test_a_webhook_that_names_a_missing_variable_stops_the_run(tmp_path: Path) -> None:
    props = '{"url": "https://example.test/h", "body": "fields:[없는것]"}'
    engine, run = start(send_process("webhook", props), env_for(tmp_path))
    assert engine.run_until_blocked(run) is State.FAILED
    assert run.error is not None and run.error.code == "webhook_var_missing"


# ─────────────────────────── 이정표 ───────────────────────────


def test_a_milestone_passes_through_and_shows_in_the_log(tmp_path: Path) -> None:
    """C14 §이벤트 — 이벤트 정의가 없는 중간 던지기가 이정표다."""
    process = straight(
        middle='<bpmn:intermediateThrowEvent id="Ms_Open" name="마감 시작"/>', node_id="Ms_Open"
    )
    engine, run = start(process, env_for(tmp_path))
    assert engine.run_until_blocked(run) is State.DONE
    states = [e for e in run.log.events if e.kind == "node_state" and e.node_id == "Ms_Open"]
    assert [e.data["task_type"] for e in states] == ["milestone", "milestone"]
    assert any("이정표: 마감 시작" == e.data.get("message") for e in run.log.events)


def test_a_signal_throw_is_not_a_milestone(tmp_path: Path) -> None:
    """신호 던지기는 이정표가 아니다 — 받는 가지가 없어도 **신호로** 다뤄진다 (조각 3d)."""
    process = straight(
        middle='<bpmn:intermediateThrowEvent id="Thr_Ready" name="ready">'
        '<bpmn:signalEventDefinition signalRef="Sig_1"/></bpmn:intermediateThrowEvent>',
        node_id="Thr_Ready",
    )
    engine, run = start(process, env_for(tmp_path))
    assert engine.run_until_blocked(run) is State.DONE
    messages = [str(e.data.get("message")) for e in run.log.events]
    assert any("신호를 보낸다" in m for m in messages)
    assert not any("이정표" in m for m in messages)


# ─────────────────────────── 실행 폴더 자체 ───────────────────────────


@pytest.mark.parametrize("raw", ["../밖.md", "/etc/passwd", "안/../../밖.md"])
def test_the_workspace_refuses_paths_outside_the_output_folder(tmp_path: Path, raw: str) -> None:
    from chaeksas.core.files import PathDenied  # noqa: PLC0415

    workspace = Workspace(output_dir=tmp_path / "outputs")
    with pytest.raises(PathDenied):
        workspace.for_write(raw)
