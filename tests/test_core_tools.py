"""AI 태스크의 내장 도구 넷 (M3) — PDF·엑셀·CSV·HTTP 읽기. ADR-0030.

거듭 보는 것 다섯.

1. **파일 도구는 실행 폴더 밖을 못 본다** (ADR-0026). 그것은 업무 실패가 **아니라** 고쳐야 할
   설정이라, `PathDenied`가 도구 경계를 그대로 넘어간다.
2. **도구는 자기 설명과 인자 모양을 들고 다닌다** — 이름만 주면 모델이 인자를 지어낸다.
3. **`http_request_tool`은 읽기만 한다** — 보내기는 `chk:send`의 일이다.
4. **자를 때는 잘랐다고 말한다** — 조용히 자르면 모델이 없는 것을 없다고 단정한다.
5. 못 읽는 것(스캔 PDF·없는 파일·깨진 형식)은 **업무 실패**다 — 오류 경계가 받을 수 있게.
"""

from __future__ import annotations

import json
import threading
from collections.abc import Iterator
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any

import pytest

from chaeksas.contracts.bpmn_ext import AiTask
from chaeksas.core.agent import tool_specs
from chaeksas.core.files import FileTaskError, PathDenied, Workspace
from chaeksas.core.tools import (
    CUT,
    HTTP_METHODS,
    MAX_ROWS,
    HttpToolError,
    ToolDef,
    builtin_tools,
    cut,
)


@pytest.fixture
def server() -> Iterator[str]:
    """127.0.0.1에 뜬 작은 서버 — 시험이 **바깥으로 나가지 않는다**."""
    bodies = {
        "/fx": (200, "application/json", json.dumps({"환율": 1387.5}, ensure_ascii=False)),
        "/text": (200, "text/plain; charset=utf-8", "오늘의 환율 고시"),
        "/down": (503, "text/plain; charset=utf-8", "쉬는 중"),
    }

    # 주소에 한글을 쓰지 않는다 — 보내는 쪽이 퍼센트 인코딩해서 `self.path`가 달라진다.
    class Handler(BaseHTTPRequestHandler):
        def _reply(self, *, body: bool) -> None:
            status, kind, text = bodies.get(self.path, (404, "text/plain", "없다"))
            raw = text.encode("utf-8")
            self.send_response(status)
            self.send_header("content-type", kind)
            self.send_header("content-length", str(len(raw)))
            self.end_headers()
            if body:
                self.wfile.write(raw)

        def do_GET(self) -> None:  # noqa: N802 — http.server가 정한 이름
            self._reply(body=True)

        def do_HEAD(self) -> None:  # noqa: N802
            self._reply(body=False)

        def log_message(self, *_: object) -> None:
            pass  # 시험 출력에 접속 기록을 쏟지 않는다

    made = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=made.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"http://127.0.0.1:{made.server_address[1]}"
    finally:
        made.shutdown()
        made.server_close()
        thread.join(timeout=2)


@pytest.fixture
def box(tmp_path: Path) -> dict[str, ToolDef]:
    """실행 폴더 하나 — 상대 경로는 **출력 폴더 기준**이다 (ADR-0026)."""
    (tmp_path / "읽는곳").mkdir()
    return builtin_tools(Workspace(output_dir=tmp_path, readable=(tmp_path,)))


def sample_pdf(target: Path, lines: list[str]) -> Path:
    """한글이 든 PDF 하나 (CID 글꼴을 심어야 진짜 문서와 같다 — 스파이크 S8)."""
    fpdf = pytest.importorskip("fpdf", reason="시험용 PDF를 만들 fpdf2가 없다")
    font = (
        Path(__file__).resolve().parent.parent
        / "packages/qt/src/chaeksas/qt/fonts/Pretendard-Regular.otf"
    )
    pdf = fpdf.FPDF()
    pdf.add_font("pretendard", "", str(font))
    pdf.set_font("pretendard", size=12)
    pdf.add_page()
    for line in lines:
        pdf.cell(0, 9, line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(target))
    return target


# ─────────────────────────── 어떤 도구가 있나 ───────────────────────────


def test_the_builtin_tools_are_there(box: dict[str, ToolDef]) -> None:
    """읽기 넷(ADR-0030) + 엑셀 쓰기 하나(ADR-0032)."""
    assert sorted(box) == [
        "csv_parser_tool",
        "excel_parser_tool",
        "excel_writer_tool",
        "http_request_tool",
        "pdf_text_tool",
    ]


def test_without_a_workspace_the_file_tools_are_absent() -> None:
    """범위 없이는 아무것도 못 읽는다 — **없다고 말하는 쪽**이 낫다 (부르면 실패하는 것보다)."""
    assert sorted(builtin_tools()) == ["http_request_tool"]


def test_a_tool_carries_its_own_description(box: dict[str, ToolDef]) -> None:
    """이름만 알려 주면 모델이 인자를 지어낸다 (ADR-0030)."""
    spec = box["csv_parser_tool"].spec
    assert spec.description
    assert "path" in spec.parameters["properties"]
    assert spec.parameters["required"] == ["path"]


def test_the_agent_hands_the_description_to_the_model(box: dict[str, ToolDef]) -> None:
    """`tool_specs()`가 이름만 넘기면 모델이 인자를 모른다."""
    task = AiTask(goal="청구서를 읽는다", domain="doc", tools=["pdf_text_tool", "없는도구"])
    (found,) = tool_specs(task, box)
    assert found.name == "pdf_text_tool"
    assert found.description and found.parameters["properties"]["path"]


def test_a_plain_function_still_works() -> None:
    """확장이 설명 없는 함수를 꽂아도 막히지 않는다 (이름만 넘어간다)."""
    task = AiTask(goal="무엇", domain="llm", tools=["손수도구"])
    (found,) = tool_specs(task, {"손수도구": lambda: "답"})
    assert found.name == "손수도구" and not found.description


# ─────────────────────────── 실행 폴더 ───────────────────────────


@pytest.mark.parametrize("tool", ["pdf_text_tool", "excel_parser_tool", "csv_parser_tool"])
def test_a_file_tool_cannot_leave_the_run_folder(box: dict[str, ToolDef], tool: str) -> None:
    """**업무 실패가 아니다** — 고쳐야 할 설정이라 `PathDenied`가 그대로 올라간다 (ADR-0026)."""
    with pytest.raises(PathDenied):
        box[tool](path="/etc/passwd")


@pytest.mark.parametrize("tool", ["pdf_text_tool", "excel_parser_tool", "csv_parser_tool"])
def test_a_missing_file_is_a_business_failure(box: dict[str, ToolDef], tmp_path: Path, tool: str) -> None:
    with pytest.raises(FileTaskError, match="파일이 없다"):
        box[tool](path="읽는곳/없는파일")


# ─────────────────────────── PDF ───────────────────────────


def test_the_pdf_tool_reads_korean(box: dict[str, ToolDef], tmp_path: Path) -> None:
    sample_pdf(tmp_path / "읽는곳" / "청구서.pdf", ["세금계산서", "합계: 1,375,000원"])
    found = box["pdf_text_tool"](path="읽는곳/청구서.pdf")
    assert "세금계산서" in found and "1,375,000원" in found


def test_the_pdf_tool_strips_the_nulls_some_producers_leave(box: dict[str, ToolDef], tmp_path: Path) -> None:
    """글자 사이에 `\\x00`이 끼어 나오는 PDF가 있다 (S8에서 실제로 봤다)."""
    sample_pdf(tmp_path / "읽는곳" / "청구서.pdf", ["세금계산서"])
    assert "\x00" not in box["pdf_text_tool"](path="읽는곳/청구서.pdf")


def test_the_pdf_tool_marks_page_numbers(box: dict[str, ToolDef], tmp_path: Path) -> None:
    sample_pdf(tmp_path / "읽는곳" / "청구서.pdf", ["첫 쪽"])
    assert "--- 1쪽 ---" in box["pdf_text_tool"](path="읽는곳/청구서.pdf")


def test_a_scanned_pdf_is_a_business_failure(box: dict[str, ToolDef], tmp_path: Path) -> None:
    """그림만 있는 PDF는 못 읽는다 — 흐름이 대처할 수 있게 **업무 실패**로 올린다."""
    sample_pdf(tmp_path / "읽는곳" / "빈것.pdf", [])
    with pytest.raises(FileTaskError, match="글자가 없다"):
        box["pdf_text_tool"](path="읽는곳/빈것.pdf")


def test_a_bad_page_range_says_so(box: dict[str, ToolDef], tmp_path: Path) -> None:
    sample_pdf(tmp_path / "읽는곳" / "청구서.pdf", ["첫 쪽"])
    with pytest.raises(FileTaskError, match="쪽 범위"):
        box["pdf_text_tool"](path="읽는곳/청구서.pdf", pages="셋")


def test_a_broken_pdf_is_a_business_failure(box: dict[str, ToolDef], tmp_path: Path) -> None:
    (tmp_path / "읽는곳" / "가짜.pdf").write_text("이건 PDF가 아니다", encoding="utf-8")
    with pytest.raises(FileTaskError):
        box["pdf_text_tool"](path="읽는곳/가짜.pdf")


# ─────────────────────────── 엑셀 ───────────────────────────


def sample_xlsx(target: Path, rows: list[list[Any]], *, sheet: str = "거래") -> Path:
    from openpyxl import Workbook

    book = Workbook()
    page = book.active
    page.title = sheet
    for row in rows:
        page.append(row)
    book.save(str(target))
    return target


def test_the_excel_tool_gives_a_json_table(box: dict[str, ToolDef], tmp_path: Path) -> None:
    """표를 **JSON 한 덩이**로 준다 — 모델이 다시 쪼개지 않아도 되게."""
    sample_xlsx(tmp_path / "읽는곳" / "거래.xlsx", [["거래처", "금액"], ["한빛상사", 1250000]])
    found = json.loads(box["excel_parser_tool"](path="읽는곳/거래.xlsx"))
    assert found["columns"] == ["거래처", "금액"]
    assert found["rows"] == [{"거래처": "한빛상사", "금액": 1250000}]
    assert found["truncated"] is False


def test_the_excel_tool_can_pick_a_sheet(box: dict[str, ToolDef], tmp_path: Path) -> None:
    from openpyxl import Workbook

    book = Workbook()
    book.active.title = "첫장"
    book.active.append(["가"])
    made = book.create_sheet("둘째장")
    made.append(["칸"])
    made.append(["값"])
    book.save(str(tmp_path / "읽는곳" / "둘.xlsx"))

    found = json.loads(box["excel_parser_tool"](path="읽는곳/둘.xlsx", sheet="둘째장"))
    assert found["sheet"] == "둘째장" and found["rows"] == [{"칸": "값"}]


def test_an_unknown_sheet_says_which_ones_exist(box: dict[str, ToolDef], tmp_path: Path) -> None:
    sample_xlsx(tmp_path / "읽는곳" / "거래.xlsx", [["가"]], sheet="거래")
    with pytest.raises(FileTaskError, match="거래"):
        box["excel_parser_tool"](path="읽는곳/거래.xlsx", sheet="없는장")


def test_the_excel_tool_says_when_it_truncated(box: dict[str, ToolDef], tmp_path: Path) -> None:
    rows: list[list[Any]] = [["번호"], *[[i] for i in range(10)]]
    sample_xlsx(tmp_path / "읽는곳" / "많다.xlsx", rows)
    found = json.loads(box["excel_parser_tool"](path="읽는곳/많다.xlsx", max_rows=3))
    assert len(found["rows"]) == 3 and found["truncated"] is True


# ─────────────────────────── CSV ───────────────────────────


def test_the_csv_tool_gives_a_json_table(box: dict[str, ToolDef], tmp_path: Path) -> None:
    (tmp_path / "읽는곳" / "카드.csv").write_text("가맹점,금액\n한빛,1200\n", encoding="utf-8")
    found = json.loads(box["csv_parser_tool"](path="읽는곳/카드.csv"))
    assert found["columns"] == ["가맹점", "금액"] and found["rows"] == [{"가맹점": "한빛", "금액": "1200"}]


def test_the_csv_tool_eats_the_excel_bom(box: dict[str, ToolDef], tmp_path: Path) -> None:
    """Excel이 쓴 CSV는 앞에 BOM이 붙는다 — 첫 칸 이름이 깨지면 뒤가 다 어긋난다."""
    (tmp_path / "읽는곳" / "엑셀.csv").write_text("가맹점,금액\n한빛,1200\n", encoding="utf-8-sig")
    found = json.loads(box["csv_parser_tool"](path="읽는곳/엑셀.csv"))
    assert found["columns"][0] == "가맹점"


def test_the_csv_tool_finds_the_delimiter(box: dict[str, ToolDef], tmp_path: Path) -> None:
    (tmp_path / "읽는곳" / "세미.csv").write_text("가맹점;금액\n한빛;1200\n", encoding="utf-8")
    found = json.loads(box["csv_parser_tool"](path="읽는곳/세미.csv"))
    assert found["columns"] == ["가맹점", "금액"]


def test_a_non_utf8_csv_says_so(box: dict[str, ToolDef], tmp_path: Path) -> None:
    (tmp_path / "읽는곳" / "cp949.csv").write_bytes("가맹점,금액\n".encode("cp949"))
    with pytest.raises(FileTaskError, match="UTF-8"):
        box["csv_parser_tool"](path="읽는곳/cp949.csv")


def test_an_empty_csv_is_not_a_failure(box: dict[str, ToolDef], tmp_path: Path) -> None:
    """빈 표는 「없다」는 답이지 실패가 아니다 (파일 목록과 같은 규칙)."""
    (tmp_path / "읽는곳" / "빈것.csv").write_text("", encoding="utf-8")
    found = json.loads(box["csv_parser_tool"](path="읽는곳/빈것.csv"))
    assert found["rows"] == [] and found["columns"] == []


def test_the_row_limit_is_capped(box: dict[str, ToolDef], tmp_path: Path) -> None:
    """모델이 `max_rows: 999999`를 지어내도 상한이 이긴다."""
    (tmp_path / "읽는곳" / "한줄.csv").write_text("가\n나\n", encoding="utf-8")
    assert json.loads(box["csv_parser_tool"](path="읽는곳/한줄.csv", max_rows=10**9))["rows"] == [{"가": "나"}]
    assert MAX_ROWS < 10**9


# ─────────────────────────── 자르기 ───────────────────────────


def test_cutting_says_it_cut() -> None:
    assert cut("12345", 10) == "12345"
    assert cut("12345", 3).startswith("123") and CUT in cut("12345", 3)


# ─────────────────────────── HTTP ───────────────────────────


def test_only_reading_methods_are_allowed(box: dict[str, ToolDef]) -> None:
    """보내기는 `chk:send`의 일이다 — 어댑터가 막는 자리를 도구로 돌아가지 못하게 한다."""
    assert HTTP_METHODS == ("GET", "HEAD")
    with pytest.raises(HttpToolError, match="읽기만"):
        box["http_request_tool"](url="https://example.test/보내기", method="POST")


def test_a_non_http_url_is_refused(box: dict[str, ToolDef]) -> None:
    with pytest.raises(HttpToolError, match="주소가 아니다"):
        box["http_request_tool"](url="file:///etc/passwd")


def test_a_korean_header_is_refused(box: dict[str, ToolDef]) -> None:
    """헤더는 ASCII다 — 보내는 쪽 라이브러리가 요청 자체를 거부한다 (CLAUDE.md §5)."""
    with pytest.raises(HttpToolError, match="ASCII"):
        box["http_request_tool"](url="https://example.test/", headers={"X-Actor": "홍길동"})


def test_the_http_tool_reads_json(box: dict[str, ToolDef], server: str) -> None:
    found = json.loads(box["http_request_tool"](url=f"{server}/fx"))
    assert found == {"환율": 1387.5}


def test_the_http_tool_reads_text(box: dict[str, ToolDef], server: str) -> None:
    assert "환율 고시" in box["http_request_tool"](url=f"{server}/text")


def test_an_error_status_is_a_business_failure(box: dict[str, ToolDef], server: str) -> None:
    with pytest.raises(HttpToolError, match="503"):
        box["http_request_tool"](url=f"{server}/down")


def test_a_head_request_gives_the_headers(box: dict[str, ToolDef], server: str) -> None:
    found = json.loads(box["http_request_tool"](url=f"{server}/fx", method="HEAD"))
    assert found["status"] == 200 and "content-type" in found["headers"]


# ─────────────────────────── 엑셀 쓰기 (ADR-0032) ───────────────────────────


def writer(tmp_path: Path, **extra: Any) -> Any:
    """쓰기 허용 폴더가 있는 도구 한 벌."""
    space = Workspace(output_dir=tmp_path / "out", **extra)
    (tmp_path / "out").mkdir(parents=True, exist_ok=True)
    return builtin_tools(space)["excel_writer_tool"]


def read_back(path: Path) -> list[tuple[Any, ...]]:
    from openpyxl import load_workbook

    book = load_workbook(str(path))
    try:
        return [tuple(row) for row in book.worksheets[0].iter_rows(values_only=True)]
    finally:
        book.close()


def test_it_makes_the_file_with_a_header(tmp_path: Path) -> None:
    tool = writer(tmp_path)
    found = json.loads(tool.run(path="장부.xlsx", rows=[{"주문번호": "A1", "금액": 100}]))
    assert found["added"] == 1 and found["skipped"] == 0
    assert read_back(tmp_path / "out" / "장부.xlsx") == [("주문번호", "금액"), ("A1", 100)]


def test_it_appends_without_touching_what_is_there(tmp_path: Path) -> None:
    """**사람의 장부를 Bot이 덮지 않는다** — 할 수 있는 일은 끝에 붙이는 것뿐이다."""
    tool = writer(tmp_path)
    tool.run(path="장부.xlsx", rows=[{"주문번호": "A1", "금액": 100}])
    tool.run(path="장부.xlsx", rows=[{"주문번호": "A2", "금액": 200}])
    assert read_back(tmp_path / "out" / "장부.xlsx") == [
        ("주문번호", "금액"), ("A1", 100), ("A2", 200)
    ]


def test_the_key_skips_what_is_already_there(tmp_path: Path) -> None:
    """같은 것을 두 번 돌려도 쌓이지 않는다 (BX-14 「같은 날 다시 실행」)."""
    tool = writer(tmp_path)
    rows = [{"주문번호": "A1", "금액": 100}, {"주문번호": "A2", "금액": 200}]
    first = json.loads(tool.run(path="장부.xlsx", rows=rows, key="주문번호"))
    again = json.loads(tool.run(path="장부.xlsx", rows=rows, key="주문번호"))
    assert (first["added"], first["skipped"]) == (2, 0)
    assert (again["added"], again["skipped"]) == (0, 2)
    assert len(read_back(tmp_path / "out" / "장부.xlsx")) == 3


def test_an_unknown_key_column_says_so(tmp_path: Path) -> None:
    tool = writer(tmp_path)
    tool.run(path="장부.xlsx", rows=[{"주문번호": "A1"}])
    with pytest.raises(FileTaskError, match="그 칸이 없다"):
        tool.run(path="장부.xlsx", rows=[{"주문번호": "A2"}], key="없는칸")


def test_rows_as_text_are_read(tmp_path: Path) -> None:
    """모델이 글로 줄 수도 있다 — JSON이면 읽는다."""
    tool = writer(tmp_path)
    found = json.loads(tool.run(path="장부.xlsx", rows='[{"주문번호": "A1"}]'))
    assert found["added"] == 1


def test_nothing_to_append_is_said_not_crashed(tmp_path: Path) -> None:
    tool = writer(tmp_path)
    found = json.loads(tool.run(path="장부.xlsx", rows=[]))
    assert found["added"] == 0 and "붙일 줄이 없다" in found["note"]


def test_writing_outside_is_denied(tmp_path: Path) -> None:
    """**쓰기 허용 폴더 밖은 못 쓴다** — 비어 있으면 출력 폴더 안만 (ADR-0032)."""
    tool = writer(tmp_path)
    with pytest.raises(PathDenied, match="쓰기 허용 폴더 밖이다"):
        tool.run(path=str(tmp_path / "남의폴더" / "장부.xlsx"), rows=[{"가": 1}])


def test_a_writable_dir_opens_that_one_place(tmp_path: Path) -> None:
    """사람이 적은 폴더만 열린다 — 기본은 비어 있다."""
    share = tmp_path / "공유"
    share.mkdir()
    tool = writer(tmp_path, writable=(share,))
    found = json.loads(tool.run(path=str(share / "주문수집.xlsx"), rows=[{"주문번호": "A1"}]))
    assert found["added"] == 1
    assert (share / "주문수집.xlsx").is_file()


def test_a_writable_dir_can_also_be_read(tmp_path: Path) -> None:
    """쓸 수 있는데 읽지 못하면 **덧붙이기를 할 수 없다** (ADR-0032)."""
    share = tmp_path / "공유"
    share.mkdir()
    tools = builtin_tools(Workspace(output_dir=tmp_path / "out", writable=(share,)))
    (tmp_path / "out").mkdir(parents=True, exist_ok=True)
    tools["excel_writer_tool"].run(path=str(share / "장부.xlsx"), rows=[{"가": 1}])
    found = json.loads(tools["excel_parser_tool"].run(path=str(share / "장부.xlsx")))
    assert found["rows"] == [{"가": 1}]
