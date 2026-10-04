"""AI 태스크가 빌려 쓰는 **내장 도구 넷** — PDF·엑셀·CSV·HTTP 읽기.

결정은 [ADR-0030](../../../../docs/decisions/0030-ai-task-tools.md)이다.

- **도구는 자기 설명과 인자 모양을 들고 다닌다** (`ToolDef.spec`). 이름만 알려 주면 모델이
  인자를 지어낸다.
- **파일 도구는 `Workspace`를 거친다** (ADR-0026) — 출력 폴더와 읽기 허용 폴더 밖은
  `PathDenied`이고 **오류 경계로 받지 않는다** (그림·설정이 잘못된 것이다). 파일이 없거나
  형식이 아니면 **업무 실패**다.
- **`http_request_tool`은 읽기만 한다** (`GET`·`HEAD`). 보내기는 `chk:send`의 일이다 —
  어댑터가 막는 자리를 도구로 돌아가지 못하게 한다.
- **돌려주는 것은 글(str)**이고, 표는 **JSON 한 덩이**다. 길면 자르고 **잘랐다고 말한다** —
  조용히 자르면 모델이 없는 것을 없다고 단정한다.
- **값은 기록하지 않는다** (원칙 6). 여기서는 아무것도 로그에 쓰지 않는다 — 궤적은 `agent`가
  이름·인자 키·길이만 남긴다.
"""

from __future__ import annotations

import csv
import io
import json
from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass
from typing import Any

from chaeksas.core.files import FileTaskError, Workspace
from chaeksas.core.llm import ToolSpec

#: 도구 하나가 돌려주는 글의 상한 (글자). 모델 맥락을 다 먹지 않게.
MAX_CHARS = 20_000
#: 표 도구가 한 번에 읽는 줄 수 상한.
MAX_ROWS = 2_000
#: `http_request_tool`이 받는 본문 상한 (바이트).
MAX_BODY = 2_000_000
#: 읽기만 하는 도구다 (ADR-0030).
HTTP_METHODS = ("GET", "HEAD")
DEFAULT_TIMEOUT_S = 30.0
MAX_REDIRECTS = 5

CUT = "\n…(여기서 잘랐다. 더 있으면 범위를 좁혀 다시 부른다)"


def cut(text: str, limit: int = MAX_CHARS) -> str:
    """상한을 넘으면 자르고 **잘랐다고 말한다**."""
    return text if len(text) <= limit else text[:limit] + CUT


def as_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, default=str)


@dataclass(frozen=True)
class ToolDef:
    """도구 하나 — 모델에게 알려 줄 설명(`spec`)과 실제로 부를 함수(`run`).

    `Tool = Callable[..., str]`과 그대로 맞물리게 `__call__`을 둔다 (`core.agent`는 부르기만
    하고, 설명이 있으면 모델에게 넘긴다).
    """

    spec: ToolSpec
    run: Callable[..., str]

    @property
    def name(self) -> str:
        return self.spec.name

    def __call__(self, **arguments: Any) -> str:
        return self.run(**arguments)


def _schema(properties: Mapping[str, Any], required: Sequence[str] = ()) -> dict[str, Any]:
    return {"type": "object", "properties": dict(properties), "required": list(required)}


_PATH = {"type": "string", "description": "읽을 파일 경로 (BPM 프로세스가 쓰는 표기 그대로, 구분자는 `/`)"}


# ─────────────────────────── PDF ───────────────────────────


def _pdf_text(workspace: Workspace, path: str, pages: str = "") -> str:
    """PDF에서 글자를 뽑는다. `pages`는 `1-3`·`2` 같은 범위 (비우면 전부)."""
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover — 설치가 깨진 자리
        raise FileTaskError("PDF를 읽는 라이브러리가 이 PC에 없다 (pypdf)") from e

    target = workspace.for_read(path)
    if not target.is_file():
        raise FileTaskError(f"파일이 없다: {path}")
    try:
        reader = PdfReader(str(target))
        wanted = _pages(pages, len(reader.pages))
        found = [_clean(reader.pages[i].extract_text() or "") for i in wanted]
    except FileTaskError:
        raise
    except Exception as e:  # noqa: BLE001 — 깨진 PDF가 무엇을 낼지 모른다
        raise FileTaskError(f"PDF를 읽지 못했다: {path} ({type(e).__name__})") from e

    text = "\n\n".join(f"--- {i + 1}쪽 ---\n{body}" for i, body in zip(wanted, found, strict=True) if body.strip())
    if not text.strip():
        # 스캔 PDF(그림만 있는 것)가 여기로 온다 — **업무 실패**다 (ADR-0030).
        raise FileTaskError(f"PDF에 글자가 없다 (스캔 문서일 수 있다): {path}")
    return cut(text)


def _clean(text: str) -> str:
    """어떤 PDF는 글자 사이에 `\\x00`이 끼어 나온다 (만든 쪽의 ToUnicode CMap 문제, S8)."""
    return text.replace("\x00", "")


def _pages(raw: str, total: int) -> list[int]:
    """`1-3`·`2`·`` → 0부터 세는 쪽 번호 목록."""
    body = raw.strip()
    if not body:
        return list(range(total))
    try:
        if "-" in body:
            start, _, end = body.partition("-")
            first, last = int(start), int(end)
        else:
            first = last = int(body)
    except ValueError as e:
        raise FileTaskError(f"쪽 범위를 읽지 못했다: {raw} (`1-3` 또는 `2`처럼 적는다)") from e
    if first < 1 or last < first:
        raise FileTaskError(f"쪽 범위가 거꾸로다: {raw}")
    return [i for i in range(first - 1, min(last, total))]


# ─────────────────────────── 엑셀 · CSV ───────────────────────────


def _excel_parser(workspace: Workspace, path: str, sheet: str = "", max_rows: int = MAX_ROWS) -> str:
    """엑셀 한 장을 **JSON 표**로 돌려준다 (`{"sheet":…, "columns":[…], "rows":[{…}]}`)."""
    from openpyxl import load_workbook

    target = workspace.for_read(path)
    if not target.is_file():
        raise FileTaskError(f"파일이 없다: {path}")
    try:
        book = load_workbook(str(target), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 — 깨진 파일이 무엇을 낼지 모른다
        raise FileTaskError(f"엑셀을 읽지 못했다: {path} ({type(e).__name__})") from e
    try:
        if sheet and sheet not in book.sheetnames:
            raise FileTaskError(f"그 시트가 없다: {sheet} (있는 것: {', '.join(book.sheetnames)})")
        page = book[sheet] if sheet else book.worksheets[0]
        # 한 줄 더 읽는다 — 그래야 「더 있었다」(`truncated`)를 말할 수 있다.
        raw = [list(row) for row in page.iter_rows(values_only=True, max_row=_rows(max_rows) + 2)]
    finally:
        book.close()
    return _table(page.title, raw, limit=_rows(max_rows))


def _csv_parser(workspace: Workspace, path: str, delimiter: str = "", max_rows: int = MAX_ROWS) -> str:
    """CSV를 **JSON 표**로 돌려준다. 구분자를 비우면 첫 줄을 보고 알아낸다."""
    target = workspace.for_read(path)
    if not target.is_file():
        raise FileTaskError(f"파일이 없다: {path}")
    try:
        # `utf-8-sig`: Excel이 쓴 CSV는 앞에 BOM이 붙는다 (첫 칸 이름이 깨진다).
        body = target.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as e:
        raise FileTaskError(f"CSV가 UTF-8이 아니다: {path}") from e
    except OSError as e:
        raise FileTaskError(f"CSV를 읽지 못했다: {path} ({e})") from e

    found = delimiter or _delimiter(body)
    rows = list(csv.reader(io.StringIO(body), delimiter=found))
    return _table(target.name, rows, limit=_rows(max_rows))


def _delimiter(body: str) -> str:
    head = body.splitlines()[0] if body.splitlines() else ""
    return max((",", ";", "\t", "|"), key=head.count) if head else ","


def _rows(raw: Any) -> int:
    try:
        found = int(raw)
    except (TypeError, ValueError):
        return MAX_ROWS
    return max(1, min(found, MAX_ROWS))


def _table(title: str, raw: Sequence[Sequence[Any]], *, limit: int) -> str:
    """첫 줄을 칸 이름으로 보고 `{이름: 값}` 목록을 만든다. 빈 표도 실패가 아니다."""
    rows = [list(row) for row in raw if any(cell not in (None, "") for cell in row)]
    if not rows:
        return as_json({"sheet": title, "columns": [], "rows": [], "truncated": False})
    columns = [str(cell) if cell not in (None, "") else f"열{i + 1}" for i, cell in enumerate(rows[0])]
    body = rows[1 : limit + 1]
    made = [{columns[i]: cell for i, cell in enumerate(row) if i < len(columns)} for row in body]
    return cut(
        as_json(
            {
                "sheet": title,
                "columns": columns,
                "rows": made,
                "truncated": len(rows) - 1 > len(made),
            }
        )
    )


# ─────────────────────────── HTTP (읽기만) ───────────────────────────


class HttpToolError(RuntimeError):
    """바깥을 읽지 못했다 — **업무 실패**다 (경계가 받는다)."""


def _http_request(url: str, method: str = "GET", headers: Mapping[str, str] | None = None) -> str:
    """주소 하나를 **읽는다**. 쓰기(POST…)는 없다 — 보내기는 `chk:send`의 일이다."""
    import httpx

    want = (method or "GET").upper()
    if want not in HTTP_METHODS:
        # 그림이 아니라 **모델**이 지어낸 것이다 — 업무 실패로 올려 흐름이 대처하게 한다.
        raise HttpToolError(f"이 도구는 읽기만 한다: {want} (되는 것: {', '.join(HTTP_METHODS)})")
    if not url.lower().startswith(("http://", "https://")):
        raise HttpToolError(f"http(s) 주소가 아니다: {url}")

    try:
        with httpx.Client(
            follow_redirects=True, max_redirects=MAX_REDIRECTS, timeout=DEFAULT_TIMEOUT_S
        ) as client:
            reply = client.request(want, url, headers=_ascii(headers))
    except httpx.HTTPError as e:
        raise HttpToolError(f"부르지 못했다: {type(e).__name__}") from e

    if reply.status_code >= 400:
        raise HttpToolError(f"{reply.status_code} 응답이다 ({url})")
    if want == "HEAD":
        return as_json({"status": reply.status_code, "headers": dict(reply.headers)})

    raw = reply.content[:MAX_BODY]
    kind = reply.headers.get("content-type", "")
    if "json" in kind:
        try:
            return cut(as_json(json.loads(raw.decode(reply.encoding or "utf-8", errors="replace"))))
        except ValueError:
            pass  # JSON이라고 했는데 아니면 글로 준다
    return cut(raw.decode(reply.encoding or "utf-8", errors="replace"))


def _ascii(headers: Mapping[str, str] | None) -> dict[str, str]:
    """**헤더 값에 한글을 넣지 않는다** (CLAUDE.md §5 — 보내는 쪽 라이브러리가 거부한다)."""
    out = {}
    for key, value in (headers or {}).items():
        text = str(value)
        if not (str(key).isascii() and text.isascii()):
            raise HttpToolError(f"헤더는 ASCII만 된다: {key}")
        out[str(key)] = text
    return out


# ─────────────────────────── 모아 주기 ───────────────────────────


def builtin_tools(workspace: Workspace | None = None) -> dict[str, ToolDef]:
    """내장 도구 넷 (ADR-0030). `workspace`가 없으면 **파일 도구는 빼고** 준다.

    빼는 이유: 파일 도구는 범위 없이는 아무것도 못 읽는다. 사전에 이름만 있고 부르면 실패하는
    것보다, **없다고 말하는 쪽**이 낫다 (모델에게도 알려 주지 않는다).
    """
    made: dict[str, ToolDef] = {
        "http_request_tool": ToolDef(
            spec=ToolSpec(
                name="http_request_tool",
                description=(
                    "주소 하나를 읽어 온다 (GET·HEAD만). JSON이면 JSON 그대로, 아니면 글로 준다. "
                    "보내거나 바꾸는 일은 할 수 없다."
                ),
                parameters=_schema(
                    {
                        "url": {"type": "string", "description": "http(s) 주소"},
                        "method": {"type": "string", "enum": list(HTTP_METHODS), "default": "GET"},
                        "headers": {"type": "object", "description": "보낼 헤더 (ASCII만)"},
                    },
                    ["url"],
                ),
            ),
            run=_http_request,
        ),
    }
    if workspace is None:
        return made

    made["pdf_text_tool"] = ToolDef(
        spec=ToolSpec(
            name="pdf_text_tool",
            description="PDF에서 글자를 뽑는다. 쪽마다 `--- N쪽 ---`을 붙인다. 스캔 문서는 글자가 없어 실패한다.",
            parameters=_schema(
                {
                    "path": _PATH,
                    "pages": {"type": "string", "description": "쪽 범위 (`1-3`·`2`). 비우면 전부"},
                },
                ["path"],
            ),
        ),
        run=lambda path, pages="": _pdf_text(workspace, path, pages),
    )
    made["excel_parser_tool"] = ToolDef(
        spec=ToolSpec(
            name="excel_parser_tool",
            description=(
                "엑셀(.xlsx) 한 장을 표로 읽는다. 첫 줄을 칸 이름으로 본다. "
                '`{"sheet":…, "columns":[…], "rows":[{…}], "truncated":…}`를 돌려준다.'
            ),
            parameters=_schema(
                {
                    "path": _PATH,
                    "sheet": {"type": "string", "description": "시트 이름. 비우면 첫 시트"},
                    "max_rows": {"type": "integer", "description": f"읽을 줄 수 (최대 {MAX_ROWS})"},
                },
                ["path"],
            ),
        ),
        run=lambda path, sheet="", max_rows=MAX_ROWS: _excel_parser(workspace, path, sheet, max_rows),
    )
    made["csv_parser_tool"] = ToolDef(
        spec=ToolSpec(
            name="csv_parser_tool",
            description=(
                "CSV를 표로 읽는다 (UTF-8). 첫 줄을 칸 이름으로 보고, 구분자를 비우면 알아서 찾는다."
            ),
            parameters=_schema(
                {
                    "path": _PATH,
                    "delimiter": {"type": "string", "description": "구분자. 비우면 자동"},
                    "max_rows": {"type": "integer", "description": f"읽을 줄 수 (최대 {MAX_ROWS})"},
                },
                ["path"],
            ),
        ),
        run=lambda path, delimiter="", max_rows=MAX_ROWS: _csv_parser(workspace, path, delimiter, max_rows),
    )
    return made


__all__ = [
    "CUT",
    "DEFAULT_TIMEOUT_S",
    "HTTP_METHODS",
    "MAX_BODY",
    "MAX_CHARS",
    "MAX_REDIRECTS",
    "MAX_ROWS",
    "HttpToolError",
    "ToolDef",
    "as_json",
    "builtin_tools",
    "cut",
]
